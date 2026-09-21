from datetime import datetime, date
from io import BytesIO
import re

import pytest
import xlwt
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app import models
from app.main import app
from app.database import SessionLocal
from app.importer import merge_import_payloads, filter_import_payload, apply_import
from app.school_importer import (ORIGINAL_BASIC_HEADERS, ORIGINAL_APPLICATION_HEADERS,
    workbook_rows, parse_school_workbook, school_template)


def xls_bytes(multiple=False, owners=False, invalid_date=False, ambiguous_exam=False):
    book = xlwt.Workbook()
    basic = book.add_sheet('改名的学生表')
    basic.write(0, 0, '基本信息')
    for i, header in enumerate(ORIGINAL_BASIC_HEADERS): basic.write(1, i, header)
    data = {'学生':'导入示例', '出生日期':datetime(2013,2,3), '入学年份':2027, '申请年级':'9年级',
        '在读学校':'示例学校', '妈妈信息（名字，电话和邮箱）':'测试家长', '学生申请邮箱':'synthetic@example.com',
        'TOEFL Junior/TOEFL账户':'SYNTHETIC-ACCOUNT', 'TOEFL Junior/TOEFL最高分':850,
        'TOEFL Junior分项分数Listening, Language, Reading':'Listening: 280, Language: 280, Reading: 290',
        'TOEFL Junior/TOEFL考试日期':'invalid' if invalid_date else datetime(2026,12,1),
        'SSAT最高分':2200, 'SSAT考试日期':datetime(2026,12,2),
        '维立克预面试日期':datetime(2026,12,3), '维立克预面试结果':'预面试示例结果',
        '维立克正式面试日期':datetime(2026,12,4), '维立克正式面试结果':'正式面试示例结果',
        'CSS完成时间':datetime(2026,11,1), 'CSS出报告时间':datetime(2026,11,5), 'CSS是否递交':'是'}
    if ambiguous_exam:
        data['TOEFL Junior分项分数Listening, Language, Reading'] = ''
    style = xlwt.easyxf(num_format_str='YYYY-MM-DD')
    for i, header in enumerate(ORIGINAL_BASIC_HEADERS):
        value = data.get(header,'')
        basic.write(2, i, value, style if isinstance(value,datetime) else xlwt.Style.default_style)
    if multiple:
        basic.write(3, ORIGINAL_BASIC_HEADERS.index('学生'), '另一示例')
        basic.write(3, ORIGINAL_BASIC_HEADERS.index('入学年份'), 2028)
    sheet = book.add_sheet('改名的申请表')
    headers = ORIGINAL_APPLICATION_HEADERS+(['学生'] if owners else [])
    for i, header in enumerate(headers): sheet.write(0, i, header)
    application = {'申请学校':'Example Academy', '所在州':'MA', '申请年级':'9年级', '截止日期':datetime(2027,1,15),
        '申请系统':'Gateway', '申请网址':'https://example.com/apply', '申请账户':'test-account', '申请密码':'synthetic-password',
        '申请费':'100 USD', 'TOEFL代码':'0123', 'TOEFL/TOEFL Junior送分日期&分数':'2026-12-10 / 850',
        'SSAT代码':'0045', 'SSAT送分日期&分数':'2026/12/11 / 2200', 'ISEE代码':'0067',
        'ISEE送分日期和分数':'2026-12-12 / 8', 'CSS':'已提交', '补充文书':'已完成', '推荐信':'两封',
        '申请特殊情况备注':'测试备注', '补交新成绩单时间':datetime(2027,2,1), '出结果时间':datetime(2027,3,10),
        '查询网址':'https://example.com/result', '账户':'test-portal', '密码':'synthetic-portal-password',
        '申请状态':'已提交', '申请结果':'等待结果', '学生':'另一示例' if multiple else '导入示例'}
    for i, header in enumerate(headers):
        value = application.get(header,'')
        sheet.write(1, i, value, style if isinstance(value,datetime) else xlwt.Style.default_style)
    output=BytesIO(); book.save(output); return output.getvalue()


def test_native_xls_all_sections_and_dates():
    payload = parse_school_workbook(workbook_rows(xls_bytes()))
    student = payload['students'][0]
    assert student['student_type'] == 'school'
    assert student['personal']['birth_date'] == '2013-02-03'
    assert student['personal']['enrollment_year'] == '2027'
    assert student['personal']['css_submitted'] == '已递交'
    assert [row['exam_name'] for row in student['exams']] == ['TOEFL Junior','SSAT']
    assert student['exams'][0]['language'] == '280'
    assert len(student['interviews']) == 2
    app_row = student['applications'][0]
    assert app_row['status'] == 'Submitted'
    assert (app_row['ssat_code'],app_row['ssat_delivery_date'],app_row['ssat_delivery_score']) == ('0045','2026-12-11','2200')
    assert app_row['application_password']=='synthetic-password' and app_row['portal_password']=='synthetic-portal-password'


def test_multi_student_requires_explicit_ownership_and_rejects_bad_dates():
    with pytest.raises(ValueError,match='归属'):
        parse_school_workbook(workbook_rows(xls_bytes(multiple=True)))
    payload = parse_school_workbook(workbook_rows(xls_bytes(multiple=True, owners=True)))
    assert not payload['students'][0]['applications']
    assert len(payload['students'][1]['applications']) == 1
    with pytest.raises(ValueError,match='有效日期'):
        parse_school_workbook(workbook_rows(xls_bytes(invalid_date=True)))
    payload = parse_school_workbook(workbook_rows(xls_bytes(ambiguous_exam=True)), 'TOEFL')
    assert payload['students'][0]['exams'][0]['exam_name']=='TOEFL' and payload['warnings']


def test_generated_xlsx_template_matches_original_columns_and_parses():
    content = school_template()
    book = load_workbook(BytesIO(content))
    assert [cell.value for cell in book['基本信息'][1]] == ORIGINAL_BASIC_HEADERS
    sheet=book['基本信息']
    sheet.cell(2,ORIGINAL_BASIC_HEADERS.index('学生')+1,'下载模板示例')
    sheet.cell(2,ORIGINAL_BASIC_HEADERS.index('入学年份')+1,2027)
    output=BytesIO();book.save(output);book.close()
    assert parse_school_workbook(workbook_rows(output.getvalue()))['students'][0]['name']=='下载模板示例'


def test_school_merge_selection_and_upsert_are_classification_safe():
    source=parse_school_workbook(workbook_rows(xls_bytes()))
    college={'students':[{'name':'导入示例','personal':{'phone':'college-only'},'applications':[],'exams':[]}]}
    merged=merge_import_payloads([source,source,college])
    assert len(merged['students'])==2 and len(merged['students'][0]['interviews'])==2
    selected=filter_import_payload(merged,['s0:student','s0:p:enrollment_year','s0:a:0','s0:e:1','s0:i:0'])
    assert len(selected['students'])==1 and selected['students'][0]['student_type']=='school'
    assert len(selected['students'][0]['interviews'])==1 and selected['students'][0]['exams'][0]['exam_name']=='SSAT'
    with SessionLocal() as db:
        apply_import(db,merged);apply_import(db,merged)
        assert db.query(models.Student).count()==2
        assert db.query(models.Application).count()==1 and db.query(models.Exam).count()==2
        assert db.query(models.ThirdPartyInterview).count()==2
        student=db.query(models.Student).filter_by(student_type='school').one()
        assert not student.phone and student.css_completed_date==date(2026,11,1)
        application=db.query(models.Application).one()
        assert application.toefl_delivery_date==date(2026,12,10) and application.result_date==date(2027,3,10)
        assert application.deadline==date(2027,1,15)


def test_upload_preview_confirm_replay_and_profile_cleanup():
    with TestClient(app) as client:
        client.headers['origin']='http://testserver'
        client.post('/login',data={'username':'admin','password':'test-only-password-93824'})
        # Parsing a bad file writes neither a preview nor student data.
        bad=client.post('/import/preview',files={'files':('wrong.xls',b'bad','application/vnd.ms-excel')})
        assert bad.status_code==400 and 'role="alert"' in bad.text
        with SessionLocal() as db:
            assert db.query(models.ImportPreview).count()==0 and db.query(models.Student).count()==0
        template=client.get('/import/template?student_type=school')
        assert template.status_code==200 and template.content.startswith(b'PK')
        response=client.post('/import/preview',files={'files':('middle high school.xls',xls_bytes(),'application/vnd.ms-excel')})
        assert response.status_code==200 and '美初美高' in response.text and '第三方面试' in response.text
        with SessionLocal() as db: assert db.query(models.Student).count()==0
        token=re.search(r'name="token" value="([a-f0-9]+)"',response.text).group(1)
        selected=re.findall(r'name="selected" value="([^"]+)" checked',response.text)
        result=client.post('/import/apply',data={'token':token,'selection_mode':'explicit','selected':selected})
        assert result.status_code==200 and '导入完成' in result.text
        assert client.post('/import/apply',data={'token':token}).status_code==410
        with SessionLocal() as db:
            student=db.query(models.Student).one();student_id=student.id;portal=student.portal_token
            student.passport_number='KEEP-HIDDEN-PASSPORT';student.target_major='KEEP-HIDDEN-MAJOR';db.commit()
        for path in [f'/students/{student_id}',f'/students/{student_id}/edit',f'/portal/{portal}/basic']:
            page=client.get(path)
            assert page.status_code==200
            for label in ['身份证号码','护照号码','信用卡','存款证明','目标方向','选课信息']:
                assert label not in page.text
            assert 'KEEP-HIDDEN' not in page.text
        html=client.get(f'/students/{student_id}').text
        assert '<details class="student-nav">' in html and '<strong>\n' not in html
        assert '<strong>-</strong>' in html
        # A school-profile save keeps unrelated legacy fields in the database.
        client.post(f'/students/{student_id}/edit',data={'name':'导入示例','student_type':'school','enrollment_year':'2027'})
        with SessionLocal() as db:
            student=db.get(models.Student,student_id)
            assert student.passport_number=='KEEP-HIDDEN-PASSPORT' and student.target_major=='KEEP-HIDDEN-MAJOR'
