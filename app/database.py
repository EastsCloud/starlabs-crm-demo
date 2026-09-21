"""PostgreSQL connections and reference data; schema changes live in migrations/."""
import os
import re
from pathlib import Path

from sqlalchemy import create_engine, inspect, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import declarative_base, sessionmaker

BASE_DIR = Path(__file__).resolve().parent.parent
SCHEMA_REVISION = '20260920_03'
DATABASE_URL = os.environ.get('DATABASE_URL', '')
if not DATABASE_URL:
    raise RuntimeError('DATABASE_URL is required; configure a PostgreSQL database.')
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL = 'postgresql://' + DATABASE_URL.split('://', 1)[1]
url = make_url(DATABASE_URL)
if url.get_backend_name() != 'postgresql':
    raise RuntimeError('DATABASE_URL must use PostgreSQL.')
url = url.set(drivername='postgresql+psycopg')
connection_args = {'connect_timeout': 10}
if os.environ.get('APP_ENV') == 'test':
    schema = os.environ.get('TEST_SCHEMA', '')
    if not re.fullmatch(r'ams_test_[a-f0-9]{32}', schema):
        raise RuntimeError('An isolated TEST_SCHEMA is required in test mode.')
    connection_args['options'] = '-csearch_path=' + schema + ' -clock_timeout=5000 -cstatement_timeout=60000'
engine = create_engine(url, connect_args=connection_args, pool_pre_ping=True, hide_parameters=True)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


def get_db():
    with SessionLocal() as db:
        yield db


def init_db():
    with engine.connect() as connection:
        if not inspect(connection).has_table('alembic_version'):
            raise RuntimeError('Database migration required: run python -m alembic upgrade head.')
        revision = connection.execute(text('SELECT version_num FROM alembic_version')).scalar()
        if revision != SCHEMA_REVISION:
            raise RuntimeError('Database schema is out of date: run python -m alembic upgrade head.')
    seed_reference_data()


REFERENCE_COLUMNS = [
    ("美国", 1, 2, 3),
    ("加州系", 4, 5, None),
    ("英国", 6, 7, None),
    ("加拿大", 8, 9, None),
    ("香港", 10, None, None),
    ("澳大利亚", 11, None, None),
    ("申请系统", 12, None, None),
    ("学校清单合集", 13, None, None),
    ("AP科目名称", 14, None, None),
]


def _cell_text(cell, pad_code=False):
    value = cell.value
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text_value = str(value).strip()
    if pad_code and text_value.isdigit():
        return text_value.zfill(4)
    return text_value


def reference_rows_from_template(path=None):
    from openpyxl import load_workbook

    template_path = Path(path) if path else BASE_DIR / "docs" / "information_template.xlsx"
    if not template_path.exists():
        return []
    workbook = load_workbook(template_path, read_only=False, data_only=True)
    try:
        if "Database" not in workbook.sheetnames:
            return [
                {"category": "美国", "name": "ACCD", "code": "4009", "secondary_code": "", "sort_order": 1, "url": "https://www.commonapp.org/"},
                {"category": "英国", "name": "Oxford", "code": "", "secondary_code": "", "sort_order": 1, "url": "https://www.ox.ac.uk/admissions/undergraduate/applying/guide-for-applicants/ucas-application"},
                {"category": "加拿大", "name": "University of Waterloo", "code": "", "secondary_code": "", "sort_order": 1, "url": "https://uwaterloo.ca/future-students/applicants/portal"},
                {"category": "申请系统", "name": "Common App", "code": "", "secondary_code": "", "sort_order": 1, "url": "https://apply.commonapp.org/"},
                {"category": "AP科目名称", "name": "AP Biology", "code": "", "secondary_code": "", "sort_order": 1, "url": ""},
            ]
        sheet = workbook["Database"]
        rows = []
        url_by_name = {}
        for _, name_col, _, _ in REFERENCE_COLUMNS:
            for row_number in range(2, sheet.max_row + 1):
                cell = sheet.cell(row_number, name_col)
                name = _cell_text(cell)
                if name and cell.hyperlink:
                    url_by_name.setdefault(name.casefold(), cell.hyperlink.target)
        fallback_urls = {
            "Leeds": "https://www.leeds.ac.uk/undergraduate-how-to-apply/doc/apply",
            "University of Leeds": "https://www.leeds.ac.uk/undergraduate-how-to-apply/doc/apply",
            "Manchester": "https://www.manchester.ac.uk/study/undergraduate/applying/how-to-apply/",
            "University of Manchester": "https://www.manchester.ac.uk/study/undergraduate/applying/how-to-apply/",
            "LSE": "https://www.lse.ac.uk/study-at-lse/Undergraduate/Prospective-Students/How-to-apply",
            "London School of Economics and Political Science": "https://www.lse.ac.uk/study-at-lse/Undergraduate/Prospective-Students/How-to-apply",
            "Warwick": "https://warwick.ac.uk/study/undergraduate/applying/how-to-apply/",
            "University of Warwick": "https://warwick.ac.uk/study/undergraduate/applying/how-to-apply/",
            "Oxford": "https://www.ox.ac.uk/admissions/undergraduate/applying/guide-for-applicants/ucas-application",
            "University of Oxford": "https://www.ox.ac.uk/admissions/undergraduate/applying/guide-for-applicants/ucas-application",
            "Imperial College London": "https://www.imperial.ac.uk/study/apply/undergraduate/",
            "University of Waterloo": "https://uwaterloo.ca/future-students/applicants/portal",
            "U of Waterloo": "https://uwaterloo.ca/future-students/applicants/portal",
            "Monash University": "https://www.monash.edu/admissions/apply/online",
        }
        for category, name_col, code_col, secondary_col in REFERENCE_COLUMNS:
            for row_number in range(2, sheet.max_row + 1):
                name_cell = sheet.cell(row_number, name_col)
                name = _cell_text(name_cell)
                if not name:
                    continue
                row_category = category
                if name_col == 4 and row_number >= 14:
                    row_category = "新加坡"
                if name == "新加坡":
                    continue
                rows.append({
                    "category": row_category,
                    "name": name,
                    "code": _cell_text(sheet.cell(row_number, code_col), True) if code_col else "",
                    "secondary_code": _cell_text(sheet.cell(row_number, secondary_col), True) if secondary_col else "",
                    "sort_order": row_number,
                    "url": (name_cell.hyperlink.target if name_cell.hyperlink else url_by_name.get(name.casefold(), fallback_urls.get(name, ""))),
                })
        return rows
    finally:
        workbook.close()


def seed_reference_data(force=False):
    from app import models

    db = SessionLocal()
    try:
        if force:
            db.query(models.ReferenceItem).delete()
            db.commit()
        elif db.query(models.ReferenceItem).count():
            return
        rows = reference_rows_from_template()
        if rows:
            db.bulk_insert_mappings(models.ReferenceItem, rows)
            db.commit()
    finally:
        db.close()
