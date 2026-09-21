from datetime import date, datetime
from io import BytesIO
import re

from openpyxl import load_workbook

from app import crud, models, school_fields


SKIP_SHEETS = {"数据统计", "模板（托福可更新新老）", "Database"}
PERSONAL_CELLS = {
    "birth_date": "A4", "phone": "A6", "address": "A8", "passport_number": "A12",
    "passport_issue_date": "A14", "passport_expiry_date": "A16", "credit_card_holder": "A18",
    "cardholder_phone": "A20", "card_channel": "A22", "card_number": "A24", "card_expiry": "A26",
    "bank_certificate_amount": "A30", "national_id_number": "A32", "chinese_address": "A34",
}
DATE_FIELDS = {"birth_date", "passport_issue_date", "passport_expiry_date", "css_completed_date", "css_report_date"}
PERSONAL_LABELS = {
    "birth_date": "出生年月日", "phone": "电话", "address": "地址", "passport_number": "护照号码",
    "passport_issue_date": "护照签发日期", "passport_expiry_date": "护照到期日期", "credit_card_holder": "信用卡卡主姓名",
    "cardholder_phone": "持卡人电话", "card_channel": "通道", "card_number": "卡号", "card_expiry": "有效期",
    "bank_certificate_amount": "存款证明金额（美元）", "national_id_number": "身份证号码", "chinese_address": "中文地址",
}
APPLICATION_PREVIEW_FIELDS = [
    ("program_name", "申请学校"), ("country", "国家"), ("batch", "批次"), ("status", "申请状态"), ("deadline", "截止时间"),
    ("form_status", "填表状态"), ("submission_date", "提交日期"), ("supplemental_essay", "补充文书"), ("transcript_required", "成绩单"),
    ("application_system", "网申系统"), ("application_username", "用户名"), ("application_password", "网申密码"),
    ("portal_id", "Portal ID"), ("portal_password", "Portal密码"), ("portal_material_progress", "Portal进度"),
    ("portal_q1", "Q1"), ("portal_a1", "A1"), ("portal_q2", "Q2"), ("portal_a2", "A2"),
    ("portal_q3", "Q3"), ("portal_a3", "A3"), ("portal_q4", "Q4"), ("portal_a4", "A4"), ("portal_q5", "Q5"), ("portal_a5", "A5"),
    ("ceeb_code", "CEEB Code"), ("language_delivery", "语言送分"), ("sat_delivery", "SAT送分"), ("act_code", "ACT Code"),
    ("act_delivery", "ACT送分"), ("ap_delivery", "AP送分"), ("other_delivery", "其他递送"), ("delivery_date", "递送日期"),
]
EXAM_PREVIEW_FIELDS = [
    ("exam_name", "名称"), ("exam_date", "日期"), ("subject", "AP科目"), ("record_locator", "SAT Record Locator"),
    ("appointment_number", "TOEFL Appointment No."), ("reading", "R/R&W"), ("listening", "L"), ("speaking", "S"),
    ("writing", "W"), ("math", "Math/M"), ("science", "Science"), ("english", "English"), ("total", "总分"),
]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    parsed = crud.parse_date(_text(value))
    return parsed.isoformat() if parsed else ""


def _country_batch(value):
    parts = [part.strip() for part in re.split(r"[\n/|]+", _text(value)) if part.strip()]
    if len(parts) >= 2:
        return parts[0], " / ".join(parts[1:])
    return (parts[0], "") if parts else ("", "")


def parse_template_workbook(content):
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    students = []
    warnings = []
    try:
        for sheet in workbook.worksheets:
            if sheet.title in SKIP_SHEETS:
                continue
            name = _text(sheet["A2"].value)
            if not name or name in {"姓名", "学生姓名"}:
                warnings.append(f"工作表 {sheet.title} 的 A2 缺少学生姓名，已跳过")
                continue
            personal = {}
            for field, cell in PERSONAL_CELLS.items():
                value = _date(sheet[cell].value) if field in DATE_FIELDS else _text(sheet[cell].value)
                if value:
                    personal[field] = value
            applications = []
            for row in range(3, 38):
                school = _text(sheet.cell(row, 3).value)
                if not school:
                    continue
                country, batch = _country_batch(sheet.cell(row, 4).value)
                qa = []
                qa_fields = {}
                for col in range(32, 42, 2):
                    question = _text(sheet.cell(row, col).value)
                    answer = _text(sheet.cell(row, col + 1).value)
                    qa_number = (col - 30) // 2
                    qa_fields[f"portal_q{qa_number}"] = question
                    qa_fields[f"portal_a{qa_number}"] = answer
                    if question or answer:
                        qa.append(f"Q: {question}\nA: {answer}")
                application = {
                    "program_name": school, "country_batch": _text(sheet.cell(row, 4).value), "country": country,
                    "batch": batch, "status": _text(sheet.cell(row, 5).value) or "WIP", "deadline": _date(sheet.cell(row, 6).value),
                    "portal_id": _text(sheet.cell(row, 7).value), "portal_password": _text(sheet.cell(row, 8).value),
                    "portal_material_progress": _text(sheet.cell(row, 9).value) or "0%", "form_status": _text(sheet.cell(row, 10).value) or "WIP",
                    "submission_date": _date(sheet.cell(row, 11).value), "supplemental_essay": _text(sheet.cell(row, 12).value),
                    "transcript_required": _text(sheet.cell(row, 13).value), "ceeb_code": _text(sheet.cell(row, 14).value),
                    "language_delivery": _text(sheet.cell(row, 15).value), "sat_delivery": _text(sheet.cell(row, 16).value),
                    "act_code": _text(sheet.cell(row, 17).value), "act_delivery": _text(sheet.cell(row, 18).value),
                    "ap_delivery": _text(sheet.cell(row, 19).value), "other_delivery": _text(sheet.cell(row, 20).value),
                    "delivery_date": _date(sheet.cell(row, 21).value), "application_system": _text(sheet.cell(row, 22).value),
                    "application_username": _text(sheet.cell(row, 23).value), "application_password": _text(sheet.cell(row, 24).value),
                    "portal_security_qa": "\n\n".join(qa),
                    "online_application_status": _text(sheet.cell(row, 10).value) or "WIP",
                    "portal_status": _text(sheet.cell(row, 9).value) or "0%",
                    "score_delivery_status": "已完成" if any(_text(sheet.cell(row, col).value) for col in [15, 16, 18, 19, 20]) else "WIP",
                }
                application.update(qa_fields)
                applications.append(application)
            exams = []
            exam_blocks = [("TOEFL", 3, 15), ("SAT", 16, 23), ("ACT", 24, 29), ("AP", 30, 38)]
            for exam_name, start, end in exam_blocks:
                for row in range(start, end):
                    exam_date = _date(sheet.cell(row, 25).value)
                    if not exam_date:
                        continue
                    record = {"exam_name": exam_name, "exam_date": exam_date, "status": "已完成"}
                    if exam_name == "TOEFL":
                        record.update(reading=_text(sheet.cell(row, 26).value), listening=_text(sheet.cell(row, 27).value), speaking=_text(sheet.cell(row, 28).value), writing=_text(sheet.cell(row, 29).value), total=_text(sheet.cell(row, 30).value), appointment_number=_text(sheet.cell(row, 31).value))
                    elif exam_name == "SAT":
                        record.update(reading=_text(sheet.cell(row, 26).value), math=_text(sheet.cell(row, 28).value), total=_text(sheet.cell(row, 30).value), record_locator=_text(sheet.cell(row, 31).value))
                    elif exam_name == "ACT":
                        record.update(math=_text(sheet.cell(row, 26).value), science=_text(sheet.cell(row, 27).value), english=_text(sheet.cell(row, 28).value), reading=_text(sheet.cell(row, 29).value), writing=_text(sheet.cell(row, 30).value), total=_text(sheet.cell(row, 31).value))
                    else:
                        record.update(subject=_text(sheet.cell(row, 26).value), total=_text(sheet.cell(row, 31).value))
                    exams.append(record)
            students.append({"name": name, "personal": personal, "applications": applications, "exams": exams})
    finally:
        workbook.close()
    return {"students": students, "warnings": warnings}


def merge_import_payloads(payloads):
    """Combine several preview payloads while keeping the first non-empty value."""
    merged = {"students": [], "warnings": [], "files": len(payloads)}
    by_name = {}
    for payload in payloads:
        merged["warnings"].extend(payload.get("warnings", []))
        for incoming in payload.get("students", []):
            identity = (incoming.get("student_type", "college"), incoming["name"])
            student = by_name.get(identity)
            if student is None:
                student = {"name": incoming["name"], "student_type": identity[0], "personal": {}, "applications": [], "exams": [], "interviews": []}
                by_name[identity] = student
                merged["students"].append(student)
            if identity[0] == 'school':
                old_year = student['personal'].get('enrollment_year')
                new_year = incoming.get('personal', {}).get('enrollment_year')
                if old_year and new_year and old_year != new_year:
                    from fastapi import HTTPException
                    raise HTTPException(409, '同名美初美高学生的入学年份不同，请核对姓名后分别导入。')
            for field, value in incoming.get("personal", {}).items():
                if _is_blank(student["personal"].get(field)) and not _is_blank(value):
                    student["personal"][field] = value
            _merge_records(student["applications"], incoming.get("applications", []), lambda row: (row.get("program_name", ""), row.get("deadline", "")))
            _merge_records(student["exams"], incoming.get("exams", []), lambda row: (row.get("exam_name", ""), row.get("exam_date", ""), row.get("subject", "")))
            _merge_records(student["interviews"], incoming.get("interviews", []), lambda row: (row.get("interview_type", ""), row.get("date", "")))
    return merged


def preview_context():
    return {
        "personal_labels": {**PERSONAL_LABELS, "enrollment_year": "入学年份", **{f.name:f.label for f in school_fields.BASIC_FIELDS}},
        "school_application_preview_fields": [(f.name, f.label) for f in school_fields.APPLICATION_FIELDS],
        "application_preview_fields": APPLICATION_PREVIEW_FIELDS,
        "exam_preview_fields": EXAM_PREVIEW_FIELDS + [("language", "Language"), ("verbal", "V"), ("quantitative", "Q"), ("analytical", "A")],
    }


def filter_import_payload(payload, selected):
    selected = set(selected)
    filtered = {"students": [], "warnings": payload.get("warnings", []), "files": payload.get("files", 1)}
    for student_index, source in enumerate(payload.get("students", [])):
        prefix = f"s{student_index}:"
        personal = {
            field: value for field, value in source.get("personal", {}).items()
            if f"{prefix}p:{field}" in selected
        }
        applications = [
            dict(row) for index, row in enumerate(source.get("applications", []))
            if f"{prefix}a:{index}" in selected
        ]
        exams = [
            dict(row) for index, row in enumerate(source.get("exams", []))
            if f"{prefix}e:{index}" in selected
        ]
        interviews = [dict(row) for index, row in enumerate(source.get("interviews", [])) if f"{prefix}i:{index}" in selected]
        if f"{prefix}student" in selected or personal or applications or exams or interviews:
            filtered["students"].append({"name": source["name"], "student_type": source.get("student_type", "college"), "personal": personal, "applications": applications, "exams": exams, "interviews": interviews})
    return filtered


def _merge_records(target, incoming, key):
    existing = {key(row): row for row in target}
    for source in incoming:
        item = existing.get(key(source))
        if item is None:
            item = dict(source)
            target.append(item)
            existing[key(item)] = item
            continue
        for field, value in source.items():
            if _is_blank(item.get(field)) and not _is_blank(value):
                item[field] = value


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def apply_import(db, payload):
    summary = {"students_created": 0, "students_updated": 0, "applications": 0, "exams": 0, "interviews": 0, "fields_filled": 0, "fields_overwritten": 0, "fields_preserved": 0}
    for row in payload.get("students", []):
        student_type = row.get("student_type", "college")
        matches = db.query(models.Student).filter(models.Student.name == row["name"], models.Student.student_type == student_type).all()
        if len(matches) > 1:
            from fastapi import HTTPException
            raise HTTPException(409, "同一分类中存在多个同名学生，请先区分姓名再导入。")
        student = matches[0] if matches else None
        student_created = student is None
        if not student:
            student = models.Student(name=row["name"], student_type=student_type, grade="11年级" if student_type == "college" else "")
            db.add(student)
            db.flush()
            summary["students_created"] += 1
        else:
            summary["students_updated"] += 1
        for field, value in row.get("personal", {}).items():
            if _is_blank(value):
                continue
            parsed_value = crud.parse_date(value) if field in DATE_FIELDS else value
            old_value = getattr(student, field, None)
            if _is_blank(old_value):
                setattr(student, field, parsed_value)
                summary["fields_filled"] += 1
            elif old_value != parsed_value:
                setattr(student, field, parsed_value)
                summary["fields_overwritten"] += 1
            else:
                summary["fields_preserved"] += 1
        for data in row.get("applications", []):
            data = dict(data)
            deadline = crud.parse_date(data.get("deadline"))
            item = db.query(models.Application).filter(
                models.Application.student_id == student.id,
                models.Application.program_name == data["program_name"],
                models.Application.deadline == deadline,
            ).first()
            if not item:
                item = db.query(models.Application).filter(
                    models.Application.student_id == student.id,
                    models.Application.program_name == data["program_name"],
                ).order_by(models.Application.id).first()
            if not item:
                item = models.Application(student_id=student.id, program_name=data["program_name"], program_type="中学" if student_type == "school" else "大学", status="WIP")
                db.add(item)
                item_is_new = True
            else:
                item_is_new = False
            for field, value in data.items():
                if field in {"program_name", "deadline"} or _is_blank(value):
                    continue
                parsed_value = crud.parse_date(value) if field in {"submission_date", "delivery_date", "result_date", "transcript_resubmit_date", "toefl_delivery_date", "ssat_delivery_date", "isee_delivery_date"} else value
                old_value = getattr(item, field, None)
                if _is_blank(old_value):
                    setattr(item, field, parsed_value)
                    summary["fields_filled"] += 1
                elif old_value != parsed_value:
                    setattr(item, field, parsed_value)
                    summary["fields_overwritten"] += 1
                else:
                    summary["fields_preserved"] += 1
            if deadline:
                if _is_blank(item.deadline):
                    item.deadline = deadline
                    summary["fields_filled"] += 1
                elif item.deadline != deadline:
                    item.deadline = deadline
                    summary["fields_overwritten"] += 1
                elif not item_is_new:
                    summary["fields_preserved"] += 1
            if item_is_new:
                item.deadline_time_node = "未添加时间信息"
            db.flush()
            summary["applications"] += 1
        for data in row.get("exams", []):
            data = dict(data)
            exam_date = crud.parse_date(data.get("exam_date"))
            item = db.query(models.Exam).filter(models.Exam.student_id == student.id, models.Exam.exam_name == data["exam_name"], models.Exam.exam_date == exam_date, models.Exam.subject == data.get("subject", "")).first()
            if not item:
                item = models.Exam(student_id=student.id, exam_name=data["exam_name"], exam_date=exam_date)
                db.add(item)
                item_is_new = True
            else:
                item_is_new = False
            for field, value in data.items():
                if field in {"exam_name", "exam_date"} or _is_blank(value):
                    continue
                old_value = getattr(item, field, None)
                if _is_blank(old_value):
                    setattr(item, field, value)
                    summary["fields_filled"] += 1
                elif old_value != value:
                    setattr(item, field, value)
                    summary["fields_overwritten"] += 1
                else:
                    summary["fields_preserved"] += 1
            if item_is_new:
                item.exam_time_node = crud.infer_time_node(student, exam_date)
            if _is_blank(item.score) and item.total:
                item.score = item.total
            db.flush()
            summary["exams"] += 1
        for data in row.get("interviews", []):
            day = crud.parse_date(data.get("date"))
            item = db.query(models.ThirdPartyInterview).filter_by(student_id=student.id, interview_type=data["interview_type"], date=day).first()
            if item is None:
                item = models.ThirdPartyInterview(student_id=student.id, interview_type=data["interview_type"], date=day)
                db.add(item)
            for field in ("result", "psee_score", "pwse_score"):
                value = data.get(field)
                if _is_blank(value):
                    continue
                old = getattr(item, field)
                key = "fields_filled" if _is_blank(old) else "fields_preserved" if old == value else "fields_overwritten"
                summary[key] += 1
                setattr(item, field, value)
            db.flush()
            summary["interviews"] += 1
    db.commit()
    return summary
