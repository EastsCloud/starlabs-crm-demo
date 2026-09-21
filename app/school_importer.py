"""Header-based school workbook import. Parsing never writes to the database."""
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import ZipFile

from openpyxl import load_workbook, Workbook
from app import school_fields

PERSONAL_HEADERS = [
    ('学生', 'name'), ('出生日期', 'birth_date'), ('出生地', 'birth_place'), ('国籍', 'nationality'),
    ('在读学校', 'current_school'), ('入学年份', 'enrollment_year'), ('申请年级', 'applying_grade'),
    ('妈妈信息（名字，电话和邮箱）', 'mother_info'), ('爸爸信息（名字，电话和邮箱）', 'father_info'),
    ('学生申请邮箱', 'application_email'), ('家长申请邮箱', 'parent_application_email'),
    ('TOEFL Junior/TOEFL账户', 'toefl_account'), ('SSAT家长账户', 'ssat_parent_account'),
    ('SSAT学生账户', 'ssat_student_account'), ('维立克账户', 'vericant_account'),
    ('Show and Tell', 'show_and_tell'), ('CSS完成时间', 'css_completed_date'),
    ('CSS出报告时间', 'css_report_date'), ('CSS是否递交', 'css_submitted'),
]
APPLICATION_HEADERS = [
    ('申请学校', 'program_name'), ('所在州', 'school_state'), ('申请年级', 'applying_grade'),
    ('截止日期', 'deadline'), ('申请系统', 'application_system'), ('申请网址', 'application_url'),
    ('申请账户', 'application_username'), ('申请密码', 'application_password'), ('申请费', 'application_fee'),
    ('TOEFL代码', 'toefl_code'), ('SSAT代码', 'ssat_code'), ('ISEE代码', 'isee_code'), ('CSS', 'css_status'),
    ('补充文书', 'school_supplemental_essay'), ('推荐信', 'recommendations'), ('申请特殊情况备注', 'notes'),
    ('补交新成绩单时间', 'transcript_resubmit_date'), ('出结果时间', 'result_date'),
    ('查询网址', 'portal_url'), ('账户', 'portal_id'), ('密码', 'portal_password'),
    ('申请状态', 'status'), ('申请结果', 'result'),
]
ORIGINAL_BASIC_HEADERS = ['序号'] + [label for label, _ in PERSONAL_HEADERS[:11]] + [
    'TOEFL Junior/TOEFL账户', 'TOEFL Junior/TOEFL最高分', 'TOEFL Junior分项分数Listening, Language, Reading',
    'TOEFL 分项分数：R L S W', 'TOEFL Junior/TOEFL考试日期', 'SSAT家长账户', 'SSAT学生账户', 'SSAT最高分',
    'SSAT考试日期', '维立克账户', '维立克预面试日期', '维立克预面试结果', '维立克正式面试日期',
    '维立克正式面试结果', 'Show and Tell', 'CSS完成时间', 'CSS出报告时间', 'CSS是否递交']
ORIGINAL_APPLICATION_HEADERS = ['序号', '申请学校', '所在州', '申请年级', '截止日期', '申请系统', '申请网址', '申请账户',
    '申请密码', '申请费', 'TOEFL代码', 'TOEFL/TOEFL Junior送分日期&分数', 'SSAT代码', 'SSAT送分日期&分数',
    'ISEE代码', 'ISEE送分日期和分数', 'CSS', '补充文书', '推荐信', '申请特殊情况备注', '补交新成绩单时间',
    '出结果时间', '查询网址', '账户', '密码', '申请状态', '申请结果']


def normalized(value):
    return re.sub(r'[\W_]+', '', str(value or '')).casefold()


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def workbook_rows(content):
    """Return native dates and strings for either XLS or XLSX, with bounded expansion."""
    if content.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        import xlrd
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            output = []
            for sheet in book.sheets():
                if sheet.nrows > 5000 or sheet.ncols > 200:
                    raise ValueError('每张表最多5000行、200列。')
                output.append((sheet.name, [[xlrd.xldate_as_datetime(cell.value, book.datemode)
                    if cell.ctype == xlrd.XL_CELL_DATE else cell.value for cell in sheet.row(r)] for r in range(sheet.nrows)]))
            return output
        finally:
            book.release_resources()
    if not content.startswith(b'PK'):
        raise ValueError('文件内容不是有效的Excel工作簿。')
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100_000_000:
            raise ValueError('工作簿解压后过大，请拆分文件。')
    book = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        output = []
        for sheet in book.worksheets:
            if (sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 200:
                raise ValueError('每张表最多5000行、200列。')
            output.append((sheet.title, list(sheet.values)))
        return output
    finally:
        book.close()


def find_header(rows, required):
    for index, row in enumerate(rows[:10]):
        columns = {normalized(value): i for i, value in enumerate(row) if cell_text(value)}
        if all(normalized(value) in columns for value in required):
            return index, columns
    return None


def is_school_workbook(sheets):
    return any(find_header(rows, ['学生', '入学年份']) for _, rows in sheets)


def date_value(value, label):
    if not cell_text(value):
        return ''
    if isinstance(value, (datetime, date)):
        return cell_text(value)
    value = cell_text(value).replace('年', '-').replace('月', '-').replace('日', '')
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f'{label}不是有效日期，请使用含年份的日期，例如2027-01-15。')


def components(value, fields, labels):
    value = cell_text(value)
    if not value:
        return {}
    numbers = re.findall(r'\d+(?:\.\d+)?', value)
    if len(numbers) != len(fields):
        raise ValueError(f'分项成绩须填写{len(fields)}个分数，顺序为{", ".join(labels)}。')
    # Accept labelled cells in any order, or plain values in the column's documented order.
    labelled = [re.search(r'(?<![A-Za-z])'+re.escape(label)+r'\s*[:：=]?\s*(\d+(?:\.\d+)?)', value, re.I) for label in labels]
    if all(labelled):
        numbers = [match.group(1) for match in labelled]
    return dict(zip(fields, numbers))


def split_delivery(value, label):
    value = cell_text(value)
    if not value:
        return '', ''
    match = re.search(r'\d{4}[-/.年]\d{1,2}[-/.月]\d{1,2}日?', value)
    if not match:
        # Preserve score/status-only cells instead of dropping an unrecognized date.
        if '/' in value or '-' in value:
            raise ValueError(f'{label}请使用“2027-01-15 / 850”格式。')
        return '', value
    delivery_date = date_value(match.group(), label)
    score = (value[:match.start()] + value[match.end():]).strip(' /&，,；;：:\n')
    score = re.sub(r'^(?:分数|成绩|送分分数)\s*[:：]?\s*', '', score)
    return delivery_date, score


def parse_school_workbook(sheets, default_exam='TOEFL Junior'):
    if default_exam not in ('TOEFL Junior', 'TOEFL'):
        raise ValueError('请选择TOEFL考试类型。')
    students, warnings, application_sheets = {}, [], []
    profile_dates = {'birth_date', 'css_completed_date', 'css_report_date'}
    for title, rows in sheets:
        header = find_header(rows, ['学生', '入学年份'])
        if not header:
            if find_header(rows, ['申请学校', '申请系统']):
                application_sheets.append((title, rows))
            continue
        index, columns = header
        for row_number, row in enumerate(rows[index+1:], index+2):
            def get(label):
                col = columns.get(normalized(label))
                return row[col] if col is not None and col < len(row) else ''
            name = cell_text(get('学生'))
            if not name:
                if any(cell_text(v) for v in row[1:]):
                    raise ValueError(f'{title}第{row_number}行缺少学生姓名。')
                continue
            if name in students:
                raise ValueError(f'{title}存在重复学生姓名，请拆分文件或合并同一学生的基本信息。')
            personal = {}
            for label, field in PERSONAL_HEADERS:
                if field == 'name':
                    continue
                value = date_value(get(label), label) if field in profile_dates else cell_text(get(label))
                if value:
                    personal[field] = value
            year = personal.get('enrollment_year', '')
            if not re.fullmatch(r'(19|20|21)\d{2}|2200', year):
                raise ValueError(f'{title}第{row_number}行请填写四位入学年份。')
            submitted = personal.get('css_submitted', '')
            if submitted:
                options = {'y': '已递交', 'yes': '已递交', '是': '已递交', '已提交': '已递交', '已递交': '已递交',
                           'n': '未递交', 'no': '未递交', '否': '未递交', '未提交': '未递交', '未递交': '未递交'}
                if submitted.casefold() not in options:
                    raise ValueError('CSS是否递交请填写是/否或已递交/未递交。')
                personal['css_submitted'] = options[submitted.casefold()]
            exams = []
            junior = components(get('TOEFL Junior分项分数Listening, Language, Reading'), ['listening','language','reading'], ['Listening','Language','Reading'])
            toefl = components(get('TOEFL 分项分数：R L S W'), ['reading','listening','speaking','writing'], ['R','L','S','W'])
            total = cell_text(get('TOEFL Junior/TOEFL最高分'))
            exam_date = date_value(get('TOEFL Junior/TOEFL考试日期'), 'TOEFL考试日期')
            if junior and toefl:
                raise ValueError('同一行只有一个TOEFL最高分/日期，请只填写对应考试的分项；另一类型请在档案标化记录中新增。')
            if junior or toefl or total or exam_date:
                exam_name = 'TOEFL Junior' if junior else 'TOEFL' if toefl else default_exam
                exams.append(dict(exam_name=exam_name, exam_date=exam_date, total=total, subject='', status='已出分' if total else '已报名', **(junior or toefl)))
                if not junior and not toefl:
                    warnings.append(f'{name}的共享TOEFL成绩按上传选项“{default_exam}”识别，请核对预览。')
            ssat_total, ssat_date = cell_text(get('SSAT最高分')), date_value(get('SSAT考试日期'), 'SSAT考试日期')
            ssat_parts = components(get('SSAT分项分数V,Q,A'), ['verbal','quantitative','analytical'], ['V','Q','A'])
            if ssat_total or ssat_date or ssat_parts:
                exams.append(dict(exam_name='SSAT', exam_date=ssat_date, total=ssat_total, subject='', **ssat_parts))
            interviews = []
            for kind in ('维立克预面试', '维立克正式面试'):
                day, result = date_value(get(kind+'日期'), kind+'日期'), cell_text(get(kind+'结果'))
                if day or result:
                    interviews.append(dict(interview_type=kind, date=day, result=result))
            students[name] = dict(name=name, student_type='school', personal=personal, applications=[], exams=exams, interviews=interviews)
    if not students:
        raise ValueError('未找到学生基本信息，请在“学生/入学年份”表头下填写数据。')
    application_dates = {field.name for field in school_fields.APPLICATION_FIELDS if field.kind == 'date'}
    for title, rows in application_sheets:
        index, columns = find_header(rows, ['申请学校', '申请系统'])
        for row_number, row in enumerate(rows[index+1:], index+2):
            def get(label):
                col = columns.get(normalized(label))
                return row[col] if col is not None and col < len(row) else ''
            if not cell_text(get('申请学校')):
                if any(cell_text(v) for v in row[1:]):
                    raise ValueError(f'{title}第{row_number}行缺少申请学校。')
                continue
            owner = cell_text(get('学生')) or cell_text(get('学生姓名'))
            if not owner and len(students) == 1:
                owner = next(iter(students))
            if owner not in students:
                raise ValueError(f'{title}第{row_number}行无法确定学生归属；多个学生时请增加“学生”列并填写基本信息表中的姓名。')
            application = {}
            for label, field in APPLICATION_HEADERS:
                value = date_value(get(label), label) if field in application_dates else cell_text(get(label))
                if value:
                    application[field] = value
            for prefix, label in [('toefl','TOEFL/TOEFL Junior送分日期&分数'), ('ssat','SSAT送分日期&分数'), ('isee','ISEE送分日期和分数')]:
                day, score = split_delivery(get(label), label)
                if day: application[prefix+'_delivery_date'] = day
                if score: application[prefix+'_delivery_score'] = score
            aliases = {'申请中':'WIP', '未开始':'WIP', '已提交':'Submitted', '已录取':'Accepted',
                       '候补':'Waitlisted', '拒绝':'Unsuccessful', '已撤回':'Withdrawn'}
            status = application.get('status')
            if status:
                application['status'] = aliases.get(status, status)
            school_fields.parse_fields({**application, 'status': application.get('status', 'WIP')}, school_fields.APPLICATION_FIELDS)
            students[owner]['applications'].append(application)
    # Fail before preview/commit if any value cannot fit the target column.
    limits = {f.name:f.limit for f in school_fields.BASIC_FIELDS}
    limits['enrollment_year'] = 4
    for student in students.values():
        if len(student['name']) > 120: raise ValueError('学生姓名过长。')
        for key, value in student['personal'].items():
            if len(value) > limits.get(key, 200): raise ValueError(f'基本信息字段{key}内容过长。')
        for application in student['applications']:
            for field in school_fields.APPLICATION_FIELDS:
                if len(application.get(field.name, '')) > field.limit: raise ValueError(f'{field.label}内容过长。')
        for exam in student['exams']:
            for field in ('total', 'reading', 'listening', 'speaking', 'writing', 'language', 'verbal', 'quantitative', 'analytical'):
                if len(exam.get(field, '')) > 30: raise ValueError('考试分数内容过长。')
        for interview in student['interviews']:
            if len(interview.get('result', '')) > 10000: raise ValueError('面试结果内容过长。')
    return {'students': list(students.values()), 'warnings': warnings}


def school_template():
    """A clean, data-free download matching the user's XLS columns; accepts XLS/XLSX."""
    book = Workbook()
    sheet = book.active; sheet.title = '基本信息'
    sheet.append(ORIGINAL_BASIC_HEADERS)
    application = book.create_sheet('申请信息'); application.append(ORIGINAL_APPLICATION_HEADERS+['学生'])
    from openpyxl.styles import Font, PatternFill, Alignment
    for sheet in book:
        sheet.freeze_panes = 'C2'; sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 48
        for cell in sheet[1]:
            cell.font = Font(name='Microsoft YaHei', color='FFFFFF', bold=True, size=11)
            cell.fill = PatternFill('solid', fgColor='2874AD'); cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.column_dimensions[cell.column_letter].width = 24
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()
